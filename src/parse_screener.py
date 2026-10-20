"""
parse_screener.py
-----------------
Turns every Screener.in "Export to Excel" file into one clean table of
company-year financials + ratios.

Input : data/raw/screener/*.xlsx   (one file per company, downloaded from Screener)
Output: data/processed/financials.csv  (one row per company per fiscal year)

Run from the project root:
    python src/parse_screener.py

All money values are in Rs crore (Screener's unit). Share price is in Rs.
"""

from pathlib import Path
import warnings

import numpy as np
import openpyxl
import pandas as pd

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

RAW = Path("data/raw/screener")
OUT = Path("data/processed")

# Screener row label -> our column name, per section of the "Data Sheet" tab
PNL_ROWS = {
    "Sales": "sales",
    "Other Income": "other_income",
    "Depreciation": "depreciation",
    "Interest": "interest",
    "Profit before tax": "pbt",
    "Tax": "tax",
    "Net profit": "net_profit",
    "Dividend Amount": "dividend",
}
BS_ROWS = {
    "Equity Share Capital": "equity_capital",
    "Reserves": "reserves",
    "Borrowings": "borrowings",
    "Other Liabilities": "other_liabilities",
    "Total": "total_assets",          # appears twice (liabilities side = assets side); we keep the first
    "Net Block": "net_block",
    "Capital Work in Progress": "cwip",
    "Investments": "investments",
    "Other Assets": "other_assets",
    "Receivables": "receivables",
    "Inventory": "inventory",
    "Cash & Bank": "cash",
    "No. of Equity Shares": "num_shares",
}
CF_ROWS = {
    "Cash from Operating Activity": "cfo",
    "Cash from Investing Activity": "cfi",
    "Cash from Financing Activity": "cff",
}
SECTION_STARTS = {
    "PROFIT & LOSS": "pnl",
    "QUARTERS": "quarters",      # quarterly data - we skip it
    "BALANCE SHEET": "bs",
    "CASH FLOW:": "cf",
    "DERIVED:": "derived",       # we skip it
}


def parse_file(path):
    """Read one Screener export and return a DataFrame with one row per fiscal year."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Data Sheet" not in wb.sheetnames:
        raise ValueError("no 'Data Sheet' tab")
    ws = wb["Data Sheet"]

    company = None
    section = None
    dates = {}        # section -> list of dates (one per column)
    years = {}        # fiscal-year date -> {column: value}

    for row in ws.iter_rows(values_only=True):
        label = row[0]
        values = list(row[1:11])
        if label is None:
            continue
        label = str(label).strip()

        if label == "COMPANY NAME":
            company = str(row[1]).strip()
            continue
        if label.upper() in SECTION_STARTS:
            section = SECTION_STARTS[label.upper()]
            continue
        if label == "Report Date":
            dates[section] = values
            continue
        if label == "PRICE:":  # year-end share price, aligned with the annual columns
            for d, v in zip(dates.get("pnl", []), values):
                if d is not None:
                    years.setdefault(d, {})["price"] = v
            continue

        mapping = {"pnl": PNL_ROWS, "bs": BS_ROWS, "cf": CF_ROWS}.get(section)
        if mapping is None or label not in mapping:
            continue
        col = mapping[label]
        for d, v in zip(dates.get(section, []), values):
            if d is None:
                continue
            rec = years.setdefault(d, {})
            if col not in rec:  # keep first occurrence (the duplicate 'Total' row)
                rec[col] = v

    wb.close()
    df = pd.DataFrame.from_dict(years, orient="index")
    df.index.name = "fy_end"
    df = df.reset_index()
    df.insert(0, "company", company)
    df.insert(1, "source_file", path.name)
    return df


def safe_div(a, b):
    out = a / b
    return out.replace([np.inf, -np.inf], np.nan)


def add_ratios(df):
    """Liquidity, leverage, profitability, cash-flow, market and trend features."""
    df = df.sort_values(["company", "fy_end"]).copy()
    num_cols = [c for c in df.columns if c not in ("company", "source_file", "fy_end")]
    df[num_cols] = df[num_cols].apply(pd.to_numeric, errors="coerce")

    df["fy_end"] = pd.to_datetime(df["fy_end"])
    df["fiscal_year"] = df["fy_end"].dt.year          # FY2017 = year ending Mar 2017
    ta = df["total_assets"]

    df["net_worth"] = df["equity_capital"] + df["reserves"]
    df["negative_net_worth"] = (df["net_worth"] <= 0).astype(int)
    df["ebit"] = df["pbt"] + df["interest"].fillna(0)

    # Leverage
    df["debt_to_equity"] = safe_div(df["borrowings"], df["net_worth"].where(df["net_worth"] > 0))
    df["debt_to_assets"] = safe_div(df["borrowings"], ta)
    df["interest_coverage"] = safe_div(df["ebit"], df["interest"].where(df["interest"] > 0))

    # Liquidity (Screener has no current-liabilities split, so this is a proxy for current ratio)
    df["liquidity_proxy"] = safe_div(df["cash"] + df["receivables"].fillna(0) + df["inventory"].fillna(0),
                                     df["other_liabilities"])
    df["cash_to_assets"] = safe_div(df["cash"], ta)

    # Profitability
    df["roa"] = safe_div(df["net_profit"], ta)
    df["ebit_margin"] = safe_div(df["ebit"], df["sales"].where(df["sales"] > 0))
    df["asset_turnover"] = safe_div(df["sales"], ta)

    # Cash flow
    df["cfo_to_assets"] = safe_div(df["cfo"], ta)
    df["cfo_to_debt"] = safe_div(df["cfo"], df["borrowings"].where(df["borrowings"] > 0))

    # Market (num_shares is an absolute count; 1 crore = 1e7)
    df["market_cap"] = df["price"] * df["num_shares"] / 1e7
    df["mcap_to_debt"] = safe_div(df["market_cap"], df["borrowings"].where(df["borrowings"] > 0))

    # Altman-style Z-score PROXY (inputs approximated from Screener fields - say so in the report)
    wc = df["cash"] + df["receivables"].fillna(0) + df["inventory"].fillna(0) - df["other_liabilities"]
    df["altman_z_proxy"] = (1.2 * safe_div(wc, ta)
                            + 1.4 * safe_div(df["reserves"], ta)
                            + 3.3 * safe_div(df["ebit"], ta)
                            + 0.6 * safe_div(df["market_cap"], df["borrowings"].fillna(0) + df["other_liabilities"])
                            + 1.0 * safe_div(df["sales"], ta))

    # Year-on-year trends within each company
    g = df.groupby("company")
    df["sales_growth"] = g["sales"].pct_change(fill_method=None).replace([np.inf, -np.inf], np.nan)
    df["borrowings_growth"] = g["borrowings"].pct_change(fill_method=None).replace([np.inf, -np.inf], np.nan)
    df["price_change"] = g["price"].pct_change(fill_method=None).replace([np.inf, -np.inf], np.nan)
    return df


def main():
    files = sorted(RAW.glob("*.xlsx"))
    if not files:
        print(f"No .xlsx files found in {RAW}")
        return

    frames, failed = [], []
    for f in files:
        try:
            frames.append(parse_file(f))
        except Exception as e:  # keep going - report bad files at the end
            failed.append((f.name, str(e)))

    if not frames:
        print("No files could be parsed:", failed)
        return

    df = add_ratios(pd.concat(frames, ignore_index=True))
    OUT.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT / "financials.csv", index=False)

    print(f"Parsed {len(frames)} of {len(files)} files -> {len(df)} company-year rows")
    print(f"Companies: {df['company'].nunique()} | fiscal years {df['fiscal_year'].min()}-{df['fiscal_year'].max()}")
    if failed:
        print("\nFailed files:")
        for name, err in failed:
            print(f"  {name}: {err}")

    show = ["company", "fiscal_year", "sales", "net_profit", "borrowings", "debt_to_equity",
            "interest_coverage", "roa", "altman_z_proxy"]
    print("\nPreview (last 6 rows):")
    print(df[show].tail(6).round(2).to_string(index=False))
    print(f"\nSaved: {OUT / 'financials.csv'}")


if __name__ == "__main__":
    main()